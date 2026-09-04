from flask import Flask, render_template, request, jsonify, send_file, flash, redirect, url_for
from models import db, Document, Sentence, Token
from config import Config
import os
import tempfile
from werkzeug.utils import secure_filename
import uuid
import shutil
from sqlalchemy import text, or_
from datetime import datetime
import nltk
from nltk.tokenize import sent_tokenize, word_tokenize
import pymorphy2
from docx import Document as DocxDocument
import PyPDF2
import re
from pathlib import Path

app = Flask(__name__)
app.config.from_object(Config)
db.init_app(app)

from corpus_service import (
    search_words, get_concordance, get_word_stats, get_lemma_stats,
    get_corpus_stats, get_document, get_documents_list, filter_by_pos
)

UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'txt', 'docx', 'pdf'}
MAX_CONTENT_LENGTH = 16 * 1024 * 1024

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = MAX_CONTENT_LENGTH

morph = pymorphy2.MorphAnalyzer()

try:
    nltk.data.find('tokenizers/punkt')
except LookupError:
    nltk.download('punkt', quiet=True)

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_txt(file_storage):
    encodings = ['utf-8', 'cp1251', 'koi8-r', 'latin1', 'cp866']
    content = file_storage.read()
    
    for enc in encodings:
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    
    return content.decode('utf-8', errors='ignore')

def extract_text_from_docx(file_storage):
    import io
    doc = DocxDocument(io.BytesIO(file_storage.read()))
    return '\n'.join([para.text for para in doc.paragraphs if para.text.strip()])

def extract_text_from_pdf(file_storage):
    import io
    text = ''
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_storage.read()))
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + '\n'
    except Exception as e:
        print(f"Ошибка чтения PDF: {e}")
    return text

def clean_text(text):
    if not text:
        return ""
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'[^\w\s.,!?;:()\-"«»]', '', text)
    return text.strip()

def process_file_direct(file_stream, filename, title, author, genre, year):
    ext = Path(filename).suffix.lower().replace('.', '')
    print(f"DEBUG process_file_direct: формат {ext}, файл {filename}")
    
    try:
        if ext == 'txt':
            encodings = ['utf-8', 'cp1251', 'koi8-r', 'latin1', 'cp866']
            content = file_stream.read()
            text = None
            
            for enc in encodings:
                try:
                    text = content.decode(enc)
                    print(f"DEBUG: TXT декодирован в {enc}")
                    break
                except UnicodeDecodeError:
                    continue
            
            if text is None:
                text = content.decode('utf-8', errors='ignore')
                print("DEBUG: TXT декодирован с игнорированием ошибок")
                
        elif ext == 'docx':
            from docx import Document as DocxDocument
            docx_doc = DocxDocument(file_stream)
            paragraphs = []
            for para in docx_doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            text = '\n'.join(paragraphs)
            print(f"DEBUG: DOCX обработан, найдено {len(paragraphs)} параграфов")
            
        elif ext == 'pdf':
            from PyPDF2 import PdfReader
            reader = PdfReader(file_stream)
            pages = []
            for page_num, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    pages.append(page_text)
            text = '\n'.join(pages)
            print(f"DEBUG: PDF обработан, {len(reader.pages)} страниц")
            
        else:
            return None, f"Неподдерживаемый формат: {ext}"
        
        if not text or len(text.strip()) < 50:
            return None, "Файл слишком мал или не содержит текста"
        
        text = clean_text(text)
        print(f"DEBUG: Текст очищен, длина: {len(text)} символов")
        
        document = Document(
            title=title,
            author=author,
            year=year,
            genre=genre,
            file_path=f"uploaded_{filename}"
        )
        db.session.add(document)
        db.session.flush()
        print(f"DEBUG: Документ создан в БД, ID: {document.id}")
        
        sent_count = 0
        token_count = 0
        
        try:
            sentences = sent_tokenize(text, language='russian')
            print(f"DEBUG: NLTK нашел {len(sentences)} предложений")
        except Exception as e:
            print(f"DEBUG: Ошибка NLTK: {e}, использую простой split")
            sentences = [s.strip() for s in text.split('.') if len(s.strip()) > 10]
        
        for sent_idx, sent_text in enumerate(sentences, start=1):
            if len(sent_text.strip()) < 10:
                continue
                
            sentence = Sentence(doc_id=document.id, text=sent_text.strip(), position=sent_idx)
            db.session.add(sentence)
            db.session.flush()
            sent_count += 1
            
            try:
                words = word_tokenize(sent_text, language='russian')
            except Exception as e:
                words = sent_text.split()
            
            words = [w for w in words if re.search(r'[а-яА-ЯёЁ]', w)]
            
            for pos_in_sent, word in enumerate(words, start=1):
                try:
                    parsed = morph.parse(word)[0]
                    lemma = parsed.normal_form
                    pos = parsed.tag.POS or 'UNKN'
                    gram = str(parsed.tag)
                    
                    token = Token(
                        sentence_id=sentence.id,
                        word=word.lower(),
                        lemma=lemma,
                        pos=pos,
                        grammemes=gram,
                        position=pos_in_sent
                    )
                    db.session.add(token)
                    token_count += 1
                    
                    if token_count % 500 == 0:
                        db.session.flush()
                        print(f"DEBUG: Обработано {token_count} токенов")
                        
                except Exception as e:
                    print(f"DEBUG: Ошибка обработки слова '{word}': {e}")
                    continue
        
        db.session.commit()
        print(f"DEBUG: Готово! {sent_count} предложений, {token_count} токенов")
        return document, f"Загружено {sent_count} предложений, {token_count} токенов"
        
    except Exception as e:
        import traceback
        print(f"DEBUG: Критическая ошибка в process_file_direct:\n{traceback.format_exc()}")
        db.session.rollback()
        return None, str(e)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/search', methods=['GET', 'POST'])
def search():
    if request.method == 'POST':
        query = request.form['query']
        search_type = request.form.get('search_type', 'word')
        results = search_words(query, search_type)
        return render_template('search.html', 
                             query=query, 
                             results=results,
                             search_type=search_type)
    return render_template('search.html')

@app.route('/search_documents', methods=['GET', 'POST'])
def search_documents():
    if request.method == 'POST':
        query = request.form.get('query', '').strip()
        search_field = request.form.get('search_field', 'all')
        
        if not query:
            return redirect(url_for('documents'))
        
        results = search_in_documents(query, search_field)
        
        return render_template('search_documents.html', 
                             query=query,
                             results=results,
                             search_field=search_field)
    
    return render_template('search_documents.html')

def search_in_documents(query, search_field='all'):
    query = f"%{query}%"
    
    if search_field == 'title':
        docs = Document.query.filter(Document.title.ilike(query)).all()
    elif search_field == 'author':
        docs = Document.query.filter(Document.author.ilike(query)).all()
    elif search_field == 'genre':
        docs = Document.query.filter(Document.genre.ilike(query)).all()
    elif search_field == 'text':
        sentences = Sentence.query.filter(Sentence.text.ilike(query)).limit(100).all()
        doc_ids = set(s.doc_id for s in sentences)
        docs = Document.query.filter(Document.id.in_(doc_ids)).all()
        
        for doc in docs:
            doc.matches = Sentence.query.filter(
                Sentence.doc_id == doc.id,
                Sentence.text.ilike(query)
            ).count()
    else:
        docs = Document.query.filter(
            or_(
                Document.title.ilike(query),
                Document.author.ilike(query),
                Document.genre.ilike(query)
            )
        ).all()
    
    return docs

@app.route('/word/<word>')
def word_detail(word):
    stats = get_word_stats(word)
    return render_template('word_detail.html', stats=stats)

@app.route('/lemma/<lemma>')
def lemma_detail(lemma):
    stats = get_lemma_stats(lemma)
    return render_template('lemma_detail.html', stats=stats)

@app.route('/concordance')
def concordance():
    word = request.args.get('word', '')
    if word:
        lines = get_concordance(word)
        return render_template('concordance.html', word=word, lines=lines)
    return render_template('concordance.html')

@app.route('/stats')
def stats():
    stats_data = get_corpus_stats()
    return render_template('stats.html', stats=stats_data)

@app.route('/documents')
def documents():
    docs = get_documents_list()
    return render_template('documents.html', documents=docs)

@app.route('/document/<int:doc_id>')
def document(doc_id):
    doc = get_document(doc_id)
    return render_template('document.html', doc=doc)

@app.route('/export/<int:doc_id>')
def export_document(doc_id):
    doc = get_document(doc_id)
    
    with tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False) as f:
        f.write(f"Название: {doc.title}\n")
        f.write(f"Автор: {doc.author}\n")
        f.write(f"Жанр: {doc.genre}\n")
        f.write(f"Год: {doc.year}\n")
        f.write("="*50 + "\n\n")
        
        for sent in doc.sentences.order_by('position'):
            f.write(sent.text + "\n")
    
    return send_file(f.name, as_attachment=True, 
                    download_name=f"{doc.title}.txt",
                    mimetype='text/plain')

@app.route('/upload', methods=['GET', 'POST'])
def upload_document():
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('Ошибка: не выбран файл')
            return redirect(request.url)
        
        file = request.files['file']
        if file.filename == '':
            flash('Ошибка: имя файла пустое')
            return redirect(request.url)
        
        original_filename = file.filename
        print(f"DEBUG: Загрузка файла: {original_filename}")
        
        if not '.' in original_filename:
            flash(f'Ошибка: файл должен иметь расширение (.txt, .docx, .pdf). Текущее имя: {original_filename}')
            return redirect(request.url)
        
        ext = original_filename.rsplit('.', 1)[1].lower()
        print(f"DEBUG: Расширение файла: {ext}")
        
        allowed = {'txt', 'docx', 'pdf'}
        if ext not in allowed:
            flash(f'Ошибка: неподдерживаемый формат "{ext}". Разрешены: txt, docx, pdf')
            return redirect(request.url)
        
        filename = secure_filename(original_filename)
        if not filename:
            filename = f"file_{uuid.uuid4().hex}.{ext}"
            print(f"DEBUG: secure_filename вернул пустое имя, генерируем: {filename}")
        
        title = request.form.get('title', '').strip()
        if not title:
            title = original_filename
            print(f"DEBUG: Используем имя файла как название: {title}")
        
        author = request.form.get('author', '').strip()
        if not author:
            author = 'Неизвестен'
        
        genre = request.form.get('genre', '').strip()
        if not genre:
            genre = 'Литературоведение'
        
        year = request.form.get('year', '').strip()
        year_int = None
        if year:
            try:
                year_int = int(year)
                if year_int < 1000 or year_int > 2100:
                    flash('Предупреждение: год вне разумного диапазона')
                    year_int = None
            except ValueError:
                flash('Предупреждение: год должен быть числом, игнорируется')
        
        print(f"DEBUG: Обработка файла: {filename}, формат: {ext}")
        print(f"DEBUG: Метаданные: {title}, {author}, {genre}, {year_int}")
        
        try:
            file_content = file.read()
            file_size = len(file_content)
            print(f"DEBUG: Размер файла: {file_size} байт")
            
            if file_size == 0:
                flash('Ошибка: файл пустой')
                return redirect(request.url)
            
            from io import BytesIO
            file_stream = BytesIO(file_content)
            
            document, message = process_file_direct(
                file_stream, 
                original_filename, 
                title, 
                author, 
                genre, 
                year_int
            )
            
            if document:
                flash(f'✅ Успех! Документ "{title}" загружен. {message}')
                return redirect(url_for('document', doc_id=document.id))
            else:
                flash(f'❌ Ошибка при обработке: {message}')
                return redirect(request.url)
                
        except Exception as e:
            import traceback
            error_trace = traceback.format_exc()
            print(f"DEBUG: Критическая ошибка при загрузке:\n{error_trace}")
            flash(f'❌ Внутренняя ошибка сервера: {str(e)}')
            return redirect(request.url)
    
    return render_template('upload.html')

@app.route('/delete_document/<int:doc_id>', methods=['POST'])
def delete_document(doc_id):
    doc = Document.query.get_or_404(doc_id)
    
    try:
        start_time = datetime.now()
        
        doc_title = doc.title
        
        db.session.execute(
            text('DELETE FROM tokens WHERE sentence_id IN (SELECT id FROM sentences WHERE doc_id = :doc_id)'),
            {'doc_id': doc_id}
        )
        
        db.session.execute(
            text('DELETE FROM sentences WHERE doc_id = :doc_id'),
            {'doc_id': doc_id}
        )
        
        db.session.execute(
            text('DELETE FROM documents WHERE id = :doc_id'),
            {'doc_id': doc_id}
        )
        
        db.session.commit()
        
        end_time = datetime.now()
        duration = (end_time - start_time).total_seconds()
        
        flash(f'Документ "{doc_title}" удален за {duration:.2f} секунд')
        
    except Exception as e:
        db.session.rollback()
        flash(f'Ошибка при удалении: {str(e)}')
    
    return redirect(url_for('documents'))

@app.route('/delete_all_documents', methods=['POST'])
def delete_all_documents():
    try:
        start_time = datetime.now()
        
        db.session.execute(text('DELETE FROM tokens'))
        db.session.execute(text('DELETE FROM sentences'))
        db.session.execute(text('DELETE FROM documents'))
        db.session.commit()
        
        end_time = datetime.now()
        duration = (end_time - start_time).total_seconds()
        
        flash(f'Все документы удалены за {duration:.2f} секунд')
        
    except Exception as e:
        db.session.rollback()
        flash(f'Ошибка: {str(e)}')
    
    return redirect(url_for('documents'))

@app.route('/filter')
def filter_page():
    return render_template('filter.html')

@app.route('/api/filter', methods=['POST'])
def api_filter():
    try:
        data = request.get_json()
        if not data:
            return jsonify({'success': False, 'error': 'Нет данных'}), 400
        
        pos = data.get('pos', '')
        if not pos:
            return jsonify({'success': False, 'error': 'Не указана часть речи'}), 400
        
        results = filter_by_pos(pos)
        
        return jsonify({
            'success': True,
            'pos': pos,
            'total': len(results),
            'results': results
        })
        
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

@app.route('/help')
def help_page():
    return render_template('help.html')

@app.route('/about')
def about():
    return render_template('about.html')

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(host='0.0.0.0', port=5000, debug=True)